"""Create a separate blank workbook for the two deferred September F4 dates.

Never overwrite an existing workbook or replace the already-filled original.
"""
from __future__ import annotations

import argparse
import sys
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

BASE = Path(__file__).resolve().parent.parent
if str(BASE) not in sys.path:
    sys.path.insert(0, str(BASE))
from tools.repair_krx_factors import atomic_write


def build_bytes() -> bytes:
    workbook = Workbook()
    guide = workbook.active
    guide.title = "입력방법"
    for text in [
        "SK스퀘어(402340) 공매도 남은 2일 입력 — 2026-10-01", "",
        "기존에 입력된 원본을 보존한 별도 템플릿입니다.",
        "KRX 포털에서 기준일 2026-09-29, 2026-09-30의 실제 당일 값을 조회하세요.",
        "잔고수량은 실제 공매도 잔고(주)이며 대차잔고로 대체하지 않습니다.",
        "잔고증감은 실제 당일 vs 직전 거래일의 값(주)입니다. 레거시 CSV로 계산하지 않습니다.",
        "잔고비율은 nav_daily.skq_shares로 자동 계산합니다. 직접 입력하지 않습니다.",
        "09-29 거래비중은 복구 완료됐습니다. 09-30 거래비중만 빈칸입니다.",
        "미공시 값은 빈칸으로 두세요. 빈칸을 0으로 바꾸거나 날짜를 이동하지 마세요.", "",
        "적재 전 CSV 사본에서 검증하고 원본 전후 SHA-256을 대조하세요.",
        "tools/apply_shorting_manual_input.py의 --workbook 인자로 이 파일을 지정하세요.",
        "--apply는 검증한 원본 해시를 --expected-sha256으로 지정한 뒤 사용하세요.",
    ]:
        guide.append([text])
    guide.column_dimensions["A"].width = 115
    sheet = workbook.create_sheet("공매도 입력")
    sheet.append(["날짜", "공매도잔고수량(주)", "공매도잔고비율(%)",
                  "당일공매도거래비중(%)", "공매도잔고증감(주, 전일대비)", "비고"])
    sheet.append(["(설명행 — 실제 데이터 아님)", None, "(자동계산)", None, None, "미공시 값은 빈칸 유지"])
    sheet.append(["2026-09-29", None, "(자동계산)", "(복구완료)", None,
                  "필요: 실제 당일 잔고수량·잔고증감"])
    sheet.append(["2026-09-30", None, "(자동계산)", None, None,
                  "필요: 실제 당일 잔고수량·잔고증감·거래비중"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    for column, width in {"A": 18, "B": 24, "C": 25, "D": 28, "E": 36, "F": 55}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "B3"
    for row in sheet.iter_rows(min_row=3, max_row=4):
        for cell in [row[1], row[3], row[4]]:
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path,
                        default=BASE / "data/manual_krx/shorting_gap_input_remaining_402340.xlsx")
    args = parser.parse_args()
    if args.output.exists():
        raise FileExistsError(f"Preserve the existing workbook: {args.output}")
    atomic_write(args.output, build_bytes())
    print("Created a blank two-date template; existing manual inputs preserved.")


if __name__ == "__main__":
    main()
